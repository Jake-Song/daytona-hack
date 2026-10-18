"""Generate the editable one-slide deck and its matching print layout.

Requires python-pptx (presentation tooling only, not an application dependency).
Run: .venv/bin/python scripts/create_pitch.py
"""
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "docs"
OUT.mkdir(exist_ok=True)
BG, PANEL, BORDER = "0B1015", "111C24", "293A46"
WHITE, MUTED, MINT, AMBER = "EFF5F8", "9DB0BE", "65DEC0", "E8BA77"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13 + 1/3), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
html = []


def unit(px):
    return Inches(px / 120)


def box(x, y, w, h, fill=PANEL, stroke=BORDER, radius=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   unit(x), unit(y), unit(w), unit(h))
    if radius:
        shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(stroke)
    shape.line.width = Pt(0.7)
    html.append(f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:#{fill};border:1px solid #{stroke};border-radius:{radius}px"></div>')


def line(x1, y1, x2, y2, color=BORDER, width=1):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, unit(x1), unit(y1), unit(x2), unit(y2))
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width * 0.6)
    html.append(f'<svg style="position:absolute;inset:0;width:1600px;height:900px" aria-hidden="true"><line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#{color}" stroke-width="{width}"/></svg>')


def text(x, y, w, h, value, size=22, color=WHITE, bold=False, link=None):
    shape = slide.shapes.add_textbox(unit(x), unit(y), unit(w), unit(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    for i, item in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.17
        run = paragraph.add_run()
        run.text = item
        run.font.name = "UnDotum"
        run.font.size = Pt(size * 0.6)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if link:
            run.hyperlink.address = link
    tag = "a" if link else "div"
    href = f' href="{escape(link)}"' if link else ""
    html.append(f'<{tag}{href} style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;font-size:{size}px;color:#{color};font-weight:{700 if bold else 400};white-space:pre;line-height:1.17;text-decoration:none">{escape(value)}</{tag}>')


# Identity / thesis.
for start, end in [((62, 46), (78, 52)), ((78, 52), (78, 66)), ((78, 66), (66, 76)), ((66, 76), (54, 66)), ((54, 66), (54, 52)), ((54, 52), (62, 46))]:
    line(*start, *end, MINT, 2)
text(96, 43, 350, 45, "AgentZero", 38, WHITE, True)
text(1130, 55, 420, 30, "NOSANA × DAYTONA  /  HACKATHON MVP", 16, MUTED)
line(54, 108, 1546, 108)
text(56, 140, 1400, 30, "SELF-HEALING SECURITY SANDBOX", 19, MINT, True)
text(54, 187, 1490, 70, "AI가 속아도, 실행 환경은 복구되어야 합니다.", 48, WHITE, True)
text(56, 264, 1480, 40, "공격을 관찰하고 → 정책을 자동 패치하고 → 실제 실행으로 방어를 검증합니다.", 25, MUTED)

# Problem.
text(56, 337, 410, 28, "01  /  THE PROBLEM", 17, MINT, True)
text(56, 387, 445, 92, "참고 문서가\n실행 지시로 바뀔 때.", 33, WHITE, True)
text(56, 500, 430, 85, "악성 README는 Agent의 행동과\n검증 범위를 왜곡할 수 있습니다.\n격리만으로 작업의 정확성은 보장되지 않습니다.", 19, MUTED)

# Architecture / control loop.
box(510, 332, 1036, 266)
text(536, 352, 700, 28, "02  /  TWO AGENTS. ONE VERIFIED CONTROL LOOP.", 17, MINT, True)
for x, title, sub, accent in [(540, "Worker Agent", "NOSANA · 작업 추론", "8DAFEA"), (870, "Security Gateway", "도구 요청 · 정책 적용", MINT), (1200, "Daytona Sandbox", "격리된 실제 코드 실행", "B6ADF1")]:
    box(x, 404, 280, 78, "16232D", BORDER, 8)
    text(x + 17, 415, 250, 28, title, 23, accent, True)
    text(x + 17, 450, 250, 22, sub, 17, MUTED)
text(830, 421, 35, 40, "→", 30, MUTED)
text(1160, 421, 35, 40, "→", 30, MUTED)
line(1338, 487, 1338, 540, MINT, 2)
line(1338, 540, 1170, 540, MINT, 2)
box(860, 511, 310, 60, "18342E", "315F50", 8)
text(880, 523, 280, 27, "Security Agent · NOSANA", 21, MINT, True)
line(1010, 505, 1010, 485, MINT, 2)
text(1021, 484, 190, 25, "↑ 정책 패치", 15, MINT)
text(548, 518, 290, 56, "실행 증거 분석 → 패치 제안\n제어 영역에서 검증·적용", 17, MUTED)

# Demonstration / customer value.
box(54, 624, 918, 186)
text(78, 644, 860, 28, "03  /  DEMO — TEST RESULT LAUNDERING", 17, MINT, True)
text(78, 690, 345, 62, "2개만 실행 → ‘전체 성공’ 보고", 24, AMBER, True)
text(463, 690, 65, 40, "→", 34, MINT)
text(548, 690, 395, 62, "3개 전체 실행 → 실패 1개 보고", 24, MINT, True)
text(80, 758, 850, 30, "복구의 목표: 실패를 숨기는 대신, 실제 테스트 결과를 정확히 드러내기.", 19, MUTED)

# Measured evidence, not invented attack success or market metrics.
box(995, 624, 551, 186)
text(1020, 644, 480, 28, "BUILT & VERIFIED", 17, MINT, True)
text(1020, 682, 100, 56, "39", 46, WHITE, True)
text(1130, 707, 130, 28, "백엔드 테스트", 18, MUTED)
text(1280, 682, 65, 56, "4", 46, WHITE, True)
text(1355, 707, 170, 28, "브라우저 테스트", 18, MUTED)
text(1020, 755, 500, 40, "실제 Daytona에서 정책 차단·전체 검증 확인", 18, MUTED)

line(54, 838, 1546, 838)
text(56, 854, 900, 28, "Worker와 분리된 제어 영역  ·  관리 키 비노출  ·  실행 증거 기반 복구", 17, MUTED)
text(1080, 853, 465, 32, "github.com/Jake-Song/daytona-hack", 19, MINT, link="https://github.com/Jake-Song/daytona-hack")

slide.notes_slide.notes_text_frame.text = (
    "30초 피치: AgentZero는 AI Agent의 행동을 관찰하고 실행 정책을 자동 복구하는 시스템입니다. "
    "Nosana는 Worker와 Security Agent의 추론을, Daytona는 격리된 실제 코드 실행을 담당합니다. "
    "대표 공격은 실패 테스트를 제외하고 전체가 성공했다고 보고하도록 유도하는 README입니다. "
    "Security Agent는 독립 검증 기록을 근거로 필수 전체 실행 정책을 제안하고, 제어 서비스가 적용·재검증합니다.\n\n"
    "검증 범위: 백엔드 39개, 브라우저 4개 테스트 통과. 실제 Daytona에서 부분 실행의 false-green과 "
    "패치 후 부분 실행·가짜 출력 차단, 전체 3개 실행에서 실제 실패 1개 검출을 확인했습니다. "
    "Qwen 실제 실행은 허위 성공을 보고하지 않아 attack_not_reproduced로 종료됐습니다. "
    "라이브 모델이 공격에 넘어간 뒤 자율 복구까지 완료했다고 주장하지 않습니다. "
    "슬라이드의 흐름은 구현된 복구 기능을 설명하며, 자동 테스트와 실제 정책 검증의 범위를 구분합니다. "
    "기존 비밀 파일 보호 시나리오도 구현되어 있습니다."
)
prs.core_properties.title = "AgentZero — Self-Healing Security Sandbox"
prs.core_properties.subject = "One-page hackathon pitch"
prs.core_properties.author = "AgentZero"
prs.save(OUT / "AgentZero-Pitch.pptx")
assert len(prs.slides) == 1
for shape in slide.shapes:
    assert shape.left >= 0 and shape.top >= 0
    assert shape.left + shape.width <= prs.slide_width + 10
    assert shape.top + shape.height <= prs.slide_height + 10
page = '<!doctype html><html lang="ko"><meta charset="UTF-8"><title>AgentZero — One-page Pitch</title><style>*{box-sizing:border-box}body{margin:0;background:#' + BG + ';font-family:UnDotum,"Noto Sans KR",Arial,sans-serif}.slide{position:relative;width:1600px;height:900px;overflow:hidden} @page{size:1600px 900px;margin:0}</style><body><main class="slide">' + ''.join(html) + '</main></body></html>'
(OUT / "AgentZero-Pitch.html").write_text(page)
print("Created one editable slide and matching print layout in docs/.")
